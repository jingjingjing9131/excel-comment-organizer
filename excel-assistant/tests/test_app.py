import base64
import json
import os
import urllib.error
import unittest
from datetime import datetime
from io import BytesIO
from unittest.mock import patch, Mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill

import app


def fixture_workbook() -> bytes:
    workbook = Workbook()
    template = workbook.active
    template.title = 'Sheet1'
    for col, heading in enumerate(app.OUTPUT_HEADERS, 1):
        template.cell(1, col, heading or None)
        template.cell(1, col).fill = PatternFill('solid', fgColor='275B47')
    for col in range(1, 14):
        template.cell(2, col).alignment = Alignment(vertical='top', wrap_text=True)

    source = workbook.create_sheet('Sheet2')
    for col, heading in enumerate(['笔记标题', '笔记链接', '内容方向', '平台', '负责人', '评论'], 1):
        source.cell(1, col, heading)
    source.append([
        '测试笔记', 'https://www.xiaohongshu.com/explore/demo', None,
        '小红书', None,
        '单评 1：第一条评论\n\n盖楼 1-1：回复评论',
    ])
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


def payload_for(data: bytes) -> dict:
    return {'workbookBase64': base64.b64encode(data).decode('ascii')}


class AppTests(unittest.TestCase):
    def setUp(self):
        self.data = fixture_workbook()

    def test_local_fallback_and_export_keep_template_and_source(self):
        payload = payload_for(self.data)
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''}):
            result = app.process_workbook(payload)
        self.assertEqual(len(result['posts']), 1)
        self.assertEqual(result['posts'][0]['splitMode'], 'local')
        self.assertEqual([c['label'] for c in result['posts'][0]['comments']], ['单评 1', '盖楼 1-1'])
        exported = app.build_export({**payload, **result})
        workbook = load_workbook(BytesIO(exported))
        self.assertIn('Sheet2原始', workbook.sheetnames)
        output = workbook['Sheet2']
        self.assertEqual(output.cell(1, 10).value, '内容端评论')
        self.assertIsNone(output.cell(1, 11).value)
        self.assertEqual(output.cell(2, 11).value, '第一条评论')
        self.assertEqual(output.cell(3, 11).value, '回复评论')
        self.assertEqual(output.cell(2, 2).value, '小红书')
        self.assertEqual(output.cell(2, 4).value, '测试笔记')
        self.assertEqual(workbook['Sheet2原始'].cell(2, 1).value, '测试笔记')
        self.assertIn('A2:A3', {str(item) for item in output.merged_cells.ranges})

    def test_deepseek_json_maps_labels_to_original_text(self):
        response = Mock()
        response.__enter__ = Mock(return_value=response)
        response.__exit__ = Mock(return_value=False)
        response.read.return_value = json.dumps({
            'choices': [{'message': {'content': json.dumps({
                'comments': [
                    {'line_indices': [1], 'label': '单评 1'},
                    {'line_indices': [2], 'label': '盖楼 1-1'},
                ]
            }, ensure_ascii=False)}}]
        }).encode('utf-8')
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-key'}), \
             patch('urllib.request.urlopen', return_value=response):
            comments = app.deepseek_split('标题', '单评1：原句一\n盖楼 1-1：原句二')
        self.assertEqual(comments, [
            {'label': '单评1', 'text': '原句一'},
            {'label': '盖楼 1-1', 'text': '原句二'},
        ])

    def test_local_split_keeps_each_reply_on_its_own_row(self):
        comments = app.local_split(
            '单评1：主评内容\n盖楼 1-1：楼层内容\n单评1回盖楼1-1：回复内容'
        )
        self.assertEqual(comments, [
            {'label': '单评1', 'text': '主评内容'},
            {'label': '盖楼 1-1', 'text': '楼层内容'},
            {'label': '单评1回盖楼1-1', 'text': '回复内容'},
        ])
        self.assertEqual(app.local_split('单评1: ：25 斤')[0]['text'], '25 斤')

    def test_deepseek_cannot_combine_several_source_lines(self):
        response = Mock()
        response.__enter__ = Mock(return_value=response)
        response.__exit__ = Mock(return_value=False)
        response.read.return_value = json.dumps({
            'choices': [{'message': {'content': json.dumps({
                'comments': [{'line_indices': [1, 2], 'label': '单评 1'}]
            }, ensure_ascii=False)}}]
        }).encode('utf-8')
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-key'}), \
             patch('urllib.request.urlopen', return_value=response):
            with self.assertRaisesRegex(app.UserError, '分组格式不正确'):
                app.deepseek_split('标题', '单评 1：主评\n盖楼 1-1：楼层')

    def test_bad_model_line_coverage_is_rejected(self):
        response = Mock()
        response.__enter__ = Mock(return_value=response)
        response.__exit__ = Mock(return_value=False)
        response.read.return_value = json.dumps({
            'choices': [{'message': {'content': '{"comments":[{"line_indices":[1],"label":"单评 1"}]}'}}]
        }).encode('utf-8')
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-key'}), \
             patch('urllib.request.urlopen', return_value=response):
            with self.assertRaises(app.UserError):
                app.deepseek_split('标题', '单评 1：原句一\n盖楼 1-1：原句二')

    def test_invalid_api_key_and_network_failure_are_reported(self):
        error_response = urllib.error.HTTPError(
            'https://api.deepseek.com/chat/completions', 401, 'Unauthorized', {},
            BytesIO(b'{"error":"invalid key"}'),
        )
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'invalid-key'}), \
             patch('urllib.request.urlopen', side_effect=error_response):
            with self.assertRaisesRegex(app.UserError, 'HTTP 401'):
                app.deepseek_split('标题', '单评 1：评论')
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': 'test-key'}), \
             patch('urllib.request.urlopen', side_effect=urllib.error.URLError('offline')):
            with self.assertRaisesRegex(app.UserError, '无法连接'):
                app.deepseek_split('标题', '单评 1：评论')

    def test_alternate_sheet_names_are_supported(self):
        workbook = load_workbook(BytesIO(self.data))
        workbook['Sheet1'].title = '模板'
        workbook['Sheet2'].title = '评论数据'
        result = BytesIO()
        workbook.save(result)
        payload = {**payload_for(result.getvalue()), 'templateSheet': '模板', 'sourceSheet': '评论数据'}
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''}):
            preview = app.process_workbook(payload)
        self.assertEqual(preview['templateSheet'], '模板')
        self.assertEqual(preview['sourceSheet'], '评论数据')

    def test_missing_source_columns_are_reported(self):
        data = fixture_workbook()
        workbook = load_workbook(BytesIO(data))
        source = workbook['Sheet2']
        source.cell(1, 6).value = '留言'
        result = BytesIO()
        workbook.save(result)
        with self.assertRaisesRegex(app.UserError, '缺少必需列'):
            app.read_posts(result.getvalue(), 'Sheet2')

    def test_selected_platform_and_calendar_date_export_as_values(self):
        payload = payload_for(self.data)
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''}):
            preview = app.process_workbook(payload)
        preview['posts'][0]['platform'] = '抖音'
        preview['posts'][0]['createdAt'] = '2026-03-04'
        exported = app.build_export({**payload, **preview})
        sheet = load_workbook(BytesIO(exported))['Sheet2']
        self.assertEqual(sheet['B2'].value, '抖音')
        self.assertEqual(sheet['F2'].value, datetime(2026, 3, 4))

    def test_invalid_calendar_date_is_rejected(self):
        payload = payload_for(self.data)
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''}):
            preview = app.process_workbook(payload)
        preview['posts'][0]['createdAt'] = 'not-a-date'
        with self.assertRaisesRegex(app.UserError, '创建时间无效'):
            app.build_export({**payload, **preview})


if __name__ == '__main__':
    unittest.main()
