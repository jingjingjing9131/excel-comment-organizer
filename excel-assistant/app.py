from __future__ import annotations

import base64
import json
import os
import re
import sys
import urllib.error
import urllib.request
from copy import copy
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
PORT = int(os.environ.get('PORT', '8765'))
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
OUTPUT_HEADERS = [
    '序号', '平台', '产品', '标题', '链接', '创建时间', '帖子负责人', '点赞截图',
    '点赞个数', '内容端评论', '', '主楼改写（尾翼）', '图片/ID',
]


def load_local_env() -> None:
    env_file = ROOT / '.env'
    if not env_file.exists():
        return
    for raw_line in env_file.read_text(encoding='utf-8').splitlines():
        line = raw_line.strip()
        if not line or line.startswith('#') or '=' not in line:
            continue
        name, value = line.split('=', 1)
        name = name.strip()
        value = value.strip().strip('"').strip("'")
        if name and name not in os.environ:
            os.environ[name] = value


load_local_env()


class UserError(Exception):
    pass


def decode_workbook(payload: dict) -> bytes:
    encoded = payload.get('workbookBase64')
    if not isinstance(encoded, str) or not encoded:
        raise UserError('请先选择一个 .xlsx 工作簿。')
    try:
        data = base64.b64decode(encoded, validate=True)
    except (ValueError, base64.binascii.Error):
        raise UserError('上传文件内容无法读取，请重新选择文件。')
    if len(data) > MAX_UPLOAD_BYTES:
        raise UserError('文件超过 20 MB，请先压缩或拆分工作簿。')
    if not data.startswith(b'PK'):
        raise UserError('文件不是有效的 .xlsx 工作簿。')
    return data


def open_uploaded_workbook(data: bytes):
    try:
        return load_workbook(BytesIO(data))
    except Exception as exc:
        raise UserError(f'无法读取工作簿：{exc}')


def header_map(sheet) -> dict[str, int]:
    return {str(cell.value).strip(): cell.column for cell in sheet[1]
            if cell.value is not None and str(cell.value).strip()}


def comment_lines(raw: str) -> list[str]:
    return [line.rstrip() for line in str(raw or '').splitlines() if line.strip()]


def infer_platform(link: str) -> str:
    link = str(link or '').lower()
    if 'xiaohongshu.com' in link:
        return '小红书'
    if 'douyin.com' in link or 'iesdouyin.com' in link:
        return '抖音'
    return ''


def extract_explicit_label(line: str) -> tuple[str, str] | None:
    match = re.match(r'^\s*([^：:]{1,100}?)\s*[：:]\s*[：:]?\s*(.*)$', line)
    if not match:
        return None
    label, text = match.group(1).strip(), match.group(2)
    if re.match(r'^(单评|主评|盖楼|贴主|作者)', label):
        return label, text
    return None


def strip_label(line: str, label: str) -> str:
    explicit = extract_explicit_label(line)
    if explicit:
        return explicit[1]
    match = re.match(r'^\s*(.*?)\s*[：:]\s*[：:]?\s*(.*)$', line)
    if not match:
        return line
    normalized_original = re.sub(r'\s+', '', match.group(1))
    normalized_model = re.sub(r'\s+', '', str(label or ''))
    return match.group(2) if normalized_model and normalized_original == normalized_model else line


def local_split(raw: str) -> list[dict[str, str]]:
    """Keep every non-empty source line as its own comment row."""
    rows = []
    for line in comment_lines(raw):
        parsed = extract_explicit_label(line)
        if parsed:
            rows.append({'label': parsed[0], 'text': parsed[1]})
        else:
            rows.append({'label': '', 'text': line})
    return rows or [{'label': '', 'text': ''}]


def deepseek_split(title: str, raw: str) -> list[dict[str, str]]:
    api_key = os.environ.get('DEEPSEEK_API_KEY', '').strip()
    if not api_key:
        raise UserError('尚未配置 DEEPSEEK_API_KEY；已改用本地拆分规则。')

    lines = comment_lines(raw)
    if not lines:
        return [{'label': '', 'text': ''}]

    model = os.environ.get('DEEPSEEK_MODEL', 'deepseek-flash').strip() or 'deepseek-flash'
    request_body = {
        'model': model,
        'temperature': 0,
        'response_format': {'type': 'json_object'},
        'messages': [
            {
                'role': 'system',
                'content': (
                    '你负责识别社交平台评论的拆分边界。输入是一篇笔记及按原顺序编号的非空原文行。'
                    '只返回 JSON：{"comments":[{"line_indices":[1],"label":"单评 1"}]}。'
                    '每一条非空原文行必须单独对应一个 comment 对象。绝对不要把多行合并成一条，'
                    '尤其是“盖楼”“单评回复”“贴主回复”等行，每行都必须单独成行。'
                    'line_indices 必须恰好只有一个编号；每行必须且只能出现一次，并按原顺序输出；编号从 1 开始。'
                    'label 只填写行首已有的评论类型和编号，去掉冒号及其后的内容；没有编号则留空。'
                    '不要改写、补充或省略评论内容。不要输出 Markdown。'
                ),
            },
            {
                'role': 'user',
                'content': json.dumps(
                    {'title': title, 'lines': [{'index': i, 'text': line}
                                               for i, line in enumerate(lines, 1)]},
                    ensure_ascii=False,
                ),
            },
        ],
    }
    request = urllib.request.Request(
        'https://api.deepseek.com/chat/completions',
        data=json.dumps(request_body).encode('utf-8'),
        headers={'Authorization': f'Bearer {api_key}', 'Content-Type': 'application/json'},
        method='POST',
    )
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            result = json.loads(response.read().decode('utf-8'))
        content = result['choices'][0]['message']['content']
        decoded = json.loads(content)
        groups = decoded.get('comments')
        if not isinstance(groups, list) or not groups:
            raise ValueError('JSON 中没有 comments 列表')

        seen: list[int] = []
        output = []
        for group in groups:
            indices = group.get('line_indices')
            label = group.get('label', '')
            if (not isinstance(indices, list) or len(indices) != 1 or
                    not isinstance(label, str) or
                    any(not isinstance(i, int) or i < 1 or i > len(lines) for i in indices)):
                raise ValueError('评论分组格式不正确')
            seen.extend(indices)
            parts = [lines[i - 1] for i in indices]
            explicit = extract_explicit_label(parts[0])
            source_label = explicit[0] if explicit else label
            parts[0] = strip_label(parts[0], label)
            output.append({'label': source_label, 'text': '\n'.join(parts)})
        if seen != list(range(1, len(lines) + 1)):
            raise ValueError('模型未能按原顺序完整且唯一地覆盖所有原文行')
        return output
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode('utf-8', errors='replace')[:400]
        raise UserError(f'DeepSeek API 返回 HTTP {exc.code}：{detail}')
    except (urllib.error.URLError, TimeoutError) as exc:
        raise UserError(f'无法连接 DeepSeek API：{exc}')
    except (KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise UserError(f'DeepSeek 返回内容无法校验：{exc}')


def read_posts(data: bytes, source_sheet_name: str) -> list[dict]:
    wb = open_uploaded_workbook(data)
    if source_sheet_name not in wb.sheetnames:
        raise UserError(f'找不到来源工作表“{source_sheet_name}”。')
    sheet = wb[source_sheet_name]
    columns = header_map(sheet)
    required = {'笔记标题', '笔记链接', '评论'}
    missing = required - columns.keys()
    if missing:
        raise UserError('来源工作表缺少必需列：' + '、'.join(sorted(missing)))

    posts = []
    for row in range(2, sheet.max_row + 1):
        title = sheet.cell(row, columns['笔记标题']).value
        link = sheet.cell(row, columns['笔记链接']).value
        raw_comments = sheet.cell(row, columns['评论']).value
        if title is None and link is None and raw_comments is None:
            continue
        link_text = str(link or '')
        source_platform = sheet.cell(row, columns['平台']).value if '平台' in columns else None
        source_platform = str(source_platform or '').strip()
        posts.append({
            'sourceRow': row,
            'title': str(title or ''),
            'link': link_text,
            'platform': source_platform if source_platform in ('抖音', '小红书') else infer_platform(link_text),
            'rawComments': str(raw_comments or ''),
        })
    if not posts:
        raise UserError('来源工作表中没有找到笔记数据。')
    return posts


def process_workbook(payload: dict) -> dict:
    data = decode_workbook(payload)
    wb = open_uploaded_workbook(data)
    names = wb.sheetnames
    template_name = payload.get('templateSheet') or ('Sheet1' if 'Sheet1' in names else '')
    source_name = payload.get('sourceSheet') or ('Sheet2' if 'Sheet2' in names else '')
    if template_name not in names or source_name not in names:
        raise UserError('请从工作表列表中分别选择格式模板和评论来源工作表。')
    template = wb[template_name]
    template_headers = [cell.value for cell in template[1][:13]]
    if [str(x or '').strip() for x in template_headers] != OUTPUT_HEADERS:
        raise UserError('模板工作表第 1 行的 13 个表头与预期格式不一致。')

    posts = read_posts(data, source_name)
    warnings = []
    for index, post in enumerate(posts, 1):
        try:
            comments = deepseek_split(post['title'], post['rawComments'])
            post['comments'] = comments
            post['splitMode'] = 'deepseek'
        except UserError as exc:
            post['comments'] = local_split(post['rawComments'])
            post['splitMode'] = 'local'
            post['warning'] = str(exc)
            # A missing key is a configuration notice, not a separate error per note.
            if not warnings or warnings[-1] != str(exc):
                warnings.append(str(exc))
        post['number'] = index
    return {
        'templateSheet': template_name,
        'sourceSheet': source_name,
        'sheetNames': names,
        'headers': OUTPUT_HEADERS,
        'posts': posts,
        'warnings': warnings,
    }


def copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)


def build_export(payload: dict) -> bytes:
    data = decode_workbook(payload)
    wb = open_uploaded_workbook(data)
    template_name = payload.get('templateSheet') or ('Sheet1' if 'Sheet1' in wb.sheetnames else '')
    source_name = payload.get('sourceSheet') or ('Sheet2' if 'Sheet2' in wb.sheetnames else '')
    if template_name not in wb.sheetnames or source_name not in wb.sheetnames:
        raise UserError('导出时找不到所选工作表，请重新上传并整理。')
    if template_name == source_name:
        raise UserError('模板工作表和来源工作表不能是同一个。')
    template = wb[template_name]
    posts = payload.get('posts')
    if not isinstance(posts, list):
        raise UserError('整理结果格式不正确，请重新处理工作簿。')

    source = wb[source_name]
    original_title = source.title
    backup_base = 'Sheet2原始'
    backup_name = backup_base
    suffix = 2
    while backup_name in wb.sheetnames and backup_name != original_title:
        backup_name = f'{backup_base}_{suffix}'
        suffix += 1
    source.title = backup_name

    output = wb.create_sheet(original_title, index=min(1, len(wb.worksheets)))
    for key, dimension in template.column_dimensions.items():
        output.column_dimensions[key].width = dimension.width
        output.column_dimensions[key].hidden = dimension.hidden
    for col in range(1, 14):
        src = template.cell(1, col)
        dst = output.cell(1, col, src.value)
        copy_cell_style(src, dst)
    output.row_dimensions[1].height = template.row_dimensions[1].height
    output.freeze_panes = template.freeze_panes

    body_style_row = 2 if template.max_row >= 2 else 1
    current_row = 2
    for post_index, post in enumerate(posts, 1):
        if not isinstance(post, dict):
            continue
        comments = post.get('comments', [])
        if not isinstance(comments, list) or not comments:
            comments = [{'label': '', 'text': post.get('rawComments', '')}]
        first = current_row
        for comment in comments:
            if not isinstance(comment, dict):
                comment = {}
            for col in range(1, 14):
                copy_cell_style(template.cell(body_style_row, col), output.cell(current_row, col))
            output.row_dimensions[current_row].height = template.row_dimensions[body_style_row].height
            output.cell(current_row, 10, str(comment.get('label') or ''))
            output.cell(current_row, 11, str(comment.get('text') or ''))
            current_row += 1
        last = current_row - 1
        link = str(post.get('link') or '')
        platform = str(post.get('platform') or infer_platform(link))
        created_at = post.get('createdAt') or None
        if created_at:
            try:
                created_at = datetime.strptime(str(created_at)[:10], '%Y-%m-%d')
            except ValueError:
                raise UserError(f'第 {post_index} 篇笔记的创建时间无效，请使用 YYYY-MM-DD 日期。')
        values = {
            1: post.get('number') or post_index,
            2: platform,
            3: post.get('product') or '',
            4: post.get('title') or '',
            5: link,
            6: created_at,
            7: post.get('owner') or '',
            8: post.get('likeScreenshot') or '',
            9: post.get('likeCount') or '',
        }
        for col, value in values.items():
            output.cell(first, col, value)
        if last > first:
            for col in range(1, 10):
                output.merge_cells(start_row=first, start_column=col, end_row=last, end_column=col)
    if current_row > 2:
        output.auto_filter.ref = f'A1:M{current_row - 1}'
    result = BytesIO()
    wb.save(result)
    return result.getvalue()


class Handler(BaseHTTPRequestHandler):
    server_version = 'ExcelOrganizer/1.0'

    def log_message(self, fmt: str, *args) -> None:
        # Keep local server output useful without logging request payloads or workbook text.
        sys.stderr.write(f'{self.address_string()} - {fmt % args}\n')

    def send_json(self, status: int, body: dict) -> None:
        encoded = json.dumps(body, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(encoded)))
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(encoded)

    def read_json(self) -> dict:
        try:
            length = int(self.headers.get('Content-Length', '0'))
            if length <= 0 or length > MAX_UPLOAD_BYTES * 2:
                raise UserError('请求体为空或超过允许大小。')
            return json.loads(self.rfile.read(length).decode('utf-8'))
        except (ValueError, UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise UserError(f'请求内容无法读取：{exc}')

    def do_GET(self) -> None:
        route = urlparse(self.path).path
        if route == '/api/status':
            self.send_json(200, {'apiKeyConfigured': bool(os.environ.get('DEEPSEEK_API_KEY', '').strip()),
                                 'model': os.environ.get('DEEPSEEK_MODEL', 'deepseek-flash')})
            return
        if route == '/app.js':
            try:
                body = (ROOT / 'app.js').read_bytes()
                self.send_response(200)
                self.send_header('Content-Type', 'text/javascript; charset=utf-8')
                self.send_header('Content-Length', str(len(body)))
                self.send_header('Cache-Control', 'no-store')
                self.end_headers()
                self.wfile.write(body)
            except OSError:
                self.send_json(500, {'error': '网站脚本文件缺失。'})
            return
        if route != '/':
            self.send_json(404, {'error': '找不到这个地址。'})
            return
        try:
            body = (ROOT / 'index.html').read_bytes()
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.send_header('Content-Length', str(len(body)))
            self.send_header('Cache-Control', 'no-store')
            self.end_headers()
            self.wfile.write(body)
        except OSError:
            self.send_json(500, {'error': '网站页面文件缺失。'})

    def do_POST(self) -> None:
        route = urlparse(self.path).path
        try:
            origin = self.headers.get('Origin')
            if origin:
                parsed_origin = urlparse(origin)
                if (parsed_origin.scheme != 'http' or
                        parsed_origin.hostname not in ('127.0.0.1', 'localhost') or
                        parsed_origin.port != self.server.server_port):
                    raise UserError('仅允许从本机整理网站发起请求。')
            payload = self.read_json()
            if route == '/api/sheets':
                data = decode_workbook(payload)
                wb = open_uploaded_workbook(data)
                self.send_json(200, {'sheetNames': wb.sheetnames,
                                     'defaultTemplate': 'Sheet1' if 'Sheet1' in wb.sheetnames else wb.sheetnames[0],
                                     'defaultSource': 'Sheet2' if 'Sheet2' in wb.sheetnames else wb.sheetnames[-1]})
                return
            if route == '/api/process':
                self.send_json(200, process_workbook(payload))
                return
            if route == '/api/export':
                result = build_export(payload)
                self.send_response(200)
                self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition', 'attachment; filename="organized.xlsx"')
                self.send_header('Content-Length', str(len(result)))
                self.send_header('Cache-Control', 'no-store')
                self.end_headers()
                self.wfile.write(result)
                return
            self.send_json(404, {'error': '找不到这个地址。'})
        except UserError as exc:
            self.send_json(400, {'error': str(exc)})
        except Exception as exc:
            self.send_json(500, {'error': f'处理失败：{exc}'})


if __name__ == '__main__':
    server = ThreadingHTTPServer(('127.0.0.1', PORT), Handler)
    print(f'Excel 整理助手已启动：http://127.0.0.1:{PORT}')
    print('按 Ctrl+C 停止服务。')
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print('\n服务已停止。')
    finally:
        server.server_close()
